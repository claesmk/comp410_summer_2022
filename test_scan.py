import unittest
import pdfplumber
import os
import re
from scan import show_aggie_pride, scan_files, get_file_text
import spacy
from presidio_analyzer import AnalyzerEngine, RecognizerRegistry, PatternRecognizer, Pattern
from docx import Document
from openpyxl import load_workbook
import pprint
pp = pprint.PrettyPrinter(indent=4)


class ScanTests(unittest.TestCase):
    # python -m spacy download en_core_web_lg
    try:
        nlp = spacy.load("en_core_web_lg")
    except OSError:
        from spacy.cli import download
        download("en_core_web_lg")
        nlp = spacy.load("en_core_web_lg")

    def test_aggie_pride(self):
        # get the slogans
        slogans = show_aggie_pride()

        self.assertIn('Aggie Pride - Worldwide', slogans)

    def test_pdf_parse(self):
        # get a platform safe path to a test pdf file
        pdf_file = os.path.join('files', 'november_statement.pdf')

        # open the file
        with pdfplumber.open(pdf_file) as pdf:
            # loop through each page in the pdf
            for p in pdf.pages:
                # extract the text as a str
                text = p.extract_text()

                # Check for the name that should be in this PDF
                self.assertIn('Name: John Smith', text)

                # Check for the account number
                self.assertIn('Account Number: 3942-29992', text)

                # Regex match on the account number
                m = re.search(r'\d+-\d+', text)
                self.assertTrue(m)

    def test_ssn(self):
        # Walk all files starting from the files directory
        for root, dirs, files in os.walk('files'):
            for name in files:
                # Look for files that should have ssn in them
                if name == 'ss_info.pdf':
                    # open the file
                    with pdfplumber.open(os.path.join(root, name)) as pdf:
                        # loop through each page in the pdf
                        for p in pdf.pages:
                            # extract the text as a str
                            text = p.extract_text()
                            # check for the ssn
                            m = re.search(r'\d{3}-\d{2}-\d{4}', text)
                            self.assertTrue(m)

    def test_scan_files(self):
        # Test to make sure scan_files returns the expected results
        expected_result = ['files/november_statement.pdf',
                           'files/Documents/twitter_info.docx',
                           'files/Documents/Statements/Retirement/ss_info.pdf',
                           'files/Downloads/address_book.xlsx',
                           'files/Downloads/address_book.txt']

        # Make expected_result os safe by checking the seperator
        if os.sep != '/':
            for i in range(len(expected_result)):
                expected_result[i] = expected_result[i].replace('/', os.sep)

        # Make sure all the expected files are actually found
        scan = scan_files()
        for f in expected_result:
            self.assertIn(f, scan)

    def test_name_recognition(self):
        # Create a sample string with a mix of things that look like they could be names
        names_string = 'Name: John Jones\nAddress: 123 John W Mitchell Drive\nNorth Carolina\nHarsh Mupali'
        doc = self.nlp(names_string)

        # expected results
        # Be careful with addresses that look like a name - they are often detected incorrectly
        expected = {'John Jones': 'PERSON',
                    '123': 'CARDINAL', 'John W Mitchell Drive': 'PERSON',
                    'North Carolina': 'GPE',
                    'Harsh Mupali': 'PERSON'}

        for ent in doc.ents:
            self.assertEqual(expected[ent.text], ent.label_)

    def test_scan_pii(self):
        # accout_recognizer = PatternRecognizer(supported_entity="US_BANK_ACCOUNT_NUMBER",
        #                                       deny_list=['3942-29992'])
        # Define the regex pattern in a Presidio `Pattern` object:
        # numbers_pattern = Pattern(name="numbers_pattern", regex="\d+", score=0.5)
        #
        # # Define the recognizer with one or more patterns
        # number_recognizer = PatternRecognizer(supported_entity="NUMBER", patterns=[numbers_pattern])

        account_pattern = Pattern(name='account_pattern', regex=r'\d{3,4}-\d{5}', score=0.9)
        account_recognizer = PatternRecognizer(supported_entity='ACCOUNT_NUMBER', patterns=[account_pattern])

        # address_recognizer = PatternRecognizer(supported_entity="STREET_ADDRESS",
        #                                        deny_list=['123 John W Mitchell Drive'])

        address_pattern = Pattern(name='address_pattern', regex=r'\d+ [A-Z][a-z]+ \w+', score=0.5)
        address_recognizer = PatternRecognizer(supported_entity='STREET_ADDRESS', patterns=[address_pattern])

        twitter_pattern = Pattern(name='twitter_pattern', regex=r'\B@\w+|at [A-Z]+\w*', score=0.5)
        twitter_recognizer = PatternRecognizer(supported_entity='TWITTER_HANDLE', patterns=[twitter_pattern])

        extra_names_pattern = Pattern(name='extra_names_pattern',
                                      regex=r'Jaden|Smith|Yuri|Krityan|Tatum', score=0.5)
        extra_names_recognizer = PatternRecognizer(supported_entity='EXTRA_PERSON', patterns=[extra_names_pattern])

        credit_card_pattern = Pattern(name='credit_card_pattern', regex=r'\d{4}[- ]\d{4}[- ]\d{4}[- ]\d{4}', score=0.9)
        cc_recognizer = PatternRecognizer(supported_entity='CREDIT_CARD_NUMBER', patterns=[credit_card_pattern])

        registry = RecognizerRegistry()
        registry.load_predefined_recognizers()

        # Add the recognizer to the existing list of recognizers
        registry.add_recognizer(account_recognizer)
        registry.add_recognizer(address_recognizer)
        registry.add_recognizer(twitter_recognizer)
        registry.add_recognizer(extra_names_recognizer)
        registry.add_recognizer(cc_recognizer)

        # Set up analyzer with our updated recognizer registry
        analyzer = AnalyzerEngine(registry=registry)

        # Detect types
        detect_types = ['ACCOUNT_NUMBER', 'STREET_ADDRESS', 'TWITTER_HANDLE', 'EXTRA_PERSON', 'CREDIT_CARD_NUMBER',
                        'PHONE_NUMBER', 'EMAIL_ADDRESS', 'US_SSN', 'PERSON']

        no_pii_list = []
        pii_list = []
        file_list = scan_files()
        for f in file_list:
            txt_list = get_file_text(f)
            for txt in txt_list:
                if txt:
                    results = analyzer.analyze(text=txt, entities=detect_types, language='en')
                    if results:
                        pii_list.append(f+':'+txt)
                        pii_list.append(results)
                    else:
                        no_pii_list.append(f+':'+txt)

        print('PII List')
        for txt in pii_list:
            print(txt)
        print('No PII list: ')
        for txt in no_pii_list:
            print(txt)

    def test_docx(self):
        # test to make sure we can read a docx ok
        doc = 'files/Documents/twitter_info.docx'

        # Fix seperator for windows (or other platforms)
        if os.sep != '/':
            doc = doc.replace('/', os.sep)

        # Read test document which contains a twitter handle
        document = Document(doc)
        for p in document.paragraphs:
            m = re.search(r'(@\w+)', p.text)
            self.assertEqual('@john_jones', m.group(1))

    def test_xlsx(self):
        xlsx = 'files/Downloads/address_book.xlsx'

        # Fix seperator for windows (or other platforms)
        if os.sep != '/':
            xlsx = xlsx.replace('/', os.sep)

        # load a workbook
        wb = load_workbook(xlsx)
        # scan through all the worksheets
        phones = []
        for ws in wb:
            for row in ws.values:
                for value in row:
                    # Find phone numbers
                    m = re.search(r'(\d{3}-\d{3}-\d{4})', value)
                    if m:
                        phones.append(m.group(1))

        # check to make sure we found all the phone numbers
        self.assertIn('336-555-1212', phones)
        self.assertIn('919-555-1212', phones)
        self.assertIn('970-555-1212', phones)

    def test_txt(self):
        txt = 'files/Downloads/address_book.txt'

        # Fix seperator for windows (or other platforms)
        if os.sep != '/':
            txt = txt.replace('/', os.sep)

        names = []
        with open(txt) as f:
            for line in f.readlines():
                m = re.search(r'([A-Z][a-z]+ [A-Z][a-z]+)', line)
                if m:
                    names.append(m.group(1))

        # check to make sure we found all the names
        self.assertIn('John Jones', names)
        self.assertIn('James Johnson', names)
        self.assertIn('Roger Jones', names)


if __name__ == '__main__':
    unittest.main()
